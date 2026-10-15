"""Motor genérico de mantenedores.

Cada mantenedor (Delegaciones, Roles, Usuarios, Ítems, Tipo y Sub Atención) es una
subclase de ``Recurso`` que declara sus campos, columnas de búsqueda y reglas de
negocio. Aquí vive todo lo común: validación, listado paginado, búsqueda por
coincidencia parcial sin distinguir mayúsculas ni tildes, exportación a Excel y
registro de rutas con control de permisos en el servidor.
"""
import io
import re
import sqlite3
from datetime import datetime
from zoneinfo import ZoneInfo

from flask import Response, g, jsonify, request
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .db import ahora_iso, get_db, normalizar
from .security import error_json, requiere
from .validators import CODIGO_RE, limpiar_texto

ZONA = ZoneInfo("America/Santiago")
TAMANOS_PAGINA = (5, 10, 25, 50, 100)


class ErrorNegocio(Exception):
    """Error controlado que se traduce en una respuesta JSON uniforme."""

    def __init__(self, status, codigo, mensaje, campos=None):
        super().__init__(mensaje)
        self.status, self.codigo, self.mensaje, self.campos = status, codigo, mensaje, campos


def fecha_local(iso):
    if not iso:
        return ""
    try:
        dt = datetime.strptime(iso, "%Y-%m-%dT%H:%M:%SZ").replace(tzinfo=ZoneInfo("UTC"))
        return dt.astimezone(ZONA).strftime("%d-%m-%Y %H:%M")
    except ValueError:
        return iso


def _seguro_excel(valor):
    """Evita inyección de fórmulas al abrir el archivo en Excel."""
    if isinstance(valor, str) and valor[:1] in ("=", "+", "-", "@", "\t", "\r"):
        return "'" + valor
    return valor


class Recurso:
    clave = None                 # prefijo de permisos y nombre del módulo
    slug = None                  # segmento de la API: /api/<slug>
    tabla = None
    etiqueta = ""
    etiquetas_estado = ("Activo", "Inactivo")
    campos = []                  # especificación de campos para validar
    busqueda = []                # expresiones SQL donde buscar
    ordenables = {}              # clave -> expresión SQL
    orden_defecto = "t.codigo"
    filtra_estado = True
    solo_lectura = False
    export_cols = []             # (encabezado, clave en el dict serializado)
    hoja = "Datos"

    # ------------------------------------------------------------ SQL base
    def desde(self):
        return f"FROM {self.tabla} t"

    def seleccion(self):
        return "t.*"

    def condiciones(self, args):
        """Alcance por delegación y filtros propios. Devuelve (where[], params[])."""
        return [], []

    def serializar(self, fila):
        d = {"id": fila["id"]}
        for c in self.campos:
            d[c["nombre"]] = fila[c["nombre"]]
        if "activo" in fila.keys():
            d["activo"] = bool(fila["activo"])
            d["estado"] = self.etiquetas_estado[0 if fila["activo"] else 1]
        return d

    # ------------------------------------------------------------- listado
    def _parametros_lista(self, args):
        try:
            pagina = max(1, int(args.get("page", 1)))
        except (TypeError, ValueError):
            pagina = 1
        try:
            tam = int(args.get("per_page", g_cfg("PAGE_SIZE_DEFAULT")))
        except (TypeError, ValueError):
            tam = g_cfg("PAGE_SIZE_DEFAULT")
        if tam not in TAMANOS_PAGINA:
            tam = g_cfg("PAGE_SIZE_DEFAULT")
        return pagina, tam

    def _where(self, args):
        where, params = self.condiciones(args)
        q = limpiar_texto(args.get("q"))[:100]
        if q and self.busqueda:
            for tok in normalizar(q).split()[:8]:
                esc = tok.replace("\\", "\\\\").replace("%", "\\%").replace("_", "\\_")
                where.append("(" + " OR ".join(f"norm({c}) LIKE ? ESCAPE '\\'" for c in self.busqueda) + ")")
                params.extend([f"%{esc}%"] * len(self.busqueda))
        estado = args.get("estado")
        if self.filtra_estado and estado in ("activo", "inactivo"):
            where.append("t.activo = ?")
            params.append(1 if estado == "activo" else 0)
        return where, params

    def _orden(self, args):
        clave = args.get("sort")
        direccion = "DESC" if args.get("dir") == "desc" else "ASC"
        expr = self.ordenables.get(clave)
        if expr:
            return f"{expr} {direccion}, t.id ASC"
        return f"{self.orden_defecto}, t.id ASC"

    def listar(self, args):
        db = get_db()
        where, params = self._where(args)
        cond = (" WHERE " + " AND ".join(where)) if where else ""
        total = db.execute(f"SELECT COUNT(*) {self.desde()}{cond}", params).fetchone()[0]
        pagina, tam = self._parametros_lista(args)
        paginas = max(1, -(-total // tam))
        pagina = min(pagina, paginas)
        filas = db.execute(
            f"SELECT {self.seleccion()} {self.desde()}{cond} ORDER BY {self._orden(args)} LIMIT ? OFFSET ?",
            params + [tam, (pagina - 1) * tam]).fetchall()
        return {
            "ok": True,
            "items": [self.serializar(f) for f in filas],
            "total": total, "page": pagina, "per_page": tam, "paginas": paginas,
        }

    def obtener(self, rid, *, exigir=True):
        """Obtiene una fila respetando el alcance del usuario (no basta con ocultar botones)."""
        db = get_db()
        where, params = self.condiciones({"_detalle": True})
        where = ["t.id = ?"] + where
        fila = db.execute(f"SELECT {self.seleccion()} {self.desde()} WHERE {' AND '.join(where)}",
                          [rid] + params).fetchone()
        if fila is None and exigir:
            existe = db.execute(f"SELECT 1 FROM {self.tabla} WHERE id=?", (rid,)).fetchone()
            if existe:
                raise ErrorNegocio(403, "fuera_de_alcance", "El registro pertenece a una delegación a la que no tienes acceso.")
            raise ErrorNegocio(404, "no_encontrado", "El registro no existe o fue eliminado.")
        return fila

    # ---------------------------------------------------------- validación
    def validar(self, datos, actual=None):
        """Validación genérica por especificación de campos + hook ``validar_extra``."""
        if not isinstance(datos, dict):
            raise ErrorNegocio(400, "solicitud_invalida", "El cuerpo de la solicitud no es válido.")
        errores, limpio = {}, {}
        for c in self.campos:
            n = c["nombre"]
            presente = n in datos
            if not presente and actual is not None:
                limpio[n] = actual[n]
                continue
            valor = datos.get(n)
            t = c.get("tipo", "texto")
            if t == "codigo":
                valor = limpiar_texto(valor).upper()
                if not valor:
                    errores[n] = f"El {c['label'].lower()} es obligatorio."
                elif not CODIGO_RE.match(valor):
                    errores[n] = "Usa de 2 a 20 caracteres: letras mayúsculas, números, guion o guion bajo."
            elif t == "choice":
                valor = limpiar_texto(valor)
                if not valor:
                    errores[n] = f"Selecciona {c['label'].lower()}."
                elif valor not in c["opciones"]:
                    errores[n] = f"{c['label']} no es una opción válida."
            else:
                valor = limpiar_texto(valor) if t == "texto" else (str(valor).strip() if valor is not None else "")
                mn, mx = c.get("min", 0), c.get("max", 255)
                if c.get("requerido", True) and not valor:
                    errores[n] = f"{c['label']} es obligatorio."
                elif valor and not (mn <= len(valor) <= mx):
                    errores[n] = f"{c['label']} debe tener entre {mn} y {mx} caracteres."
            limpio[n] = valor
        if self.filtra_estado or "activo" in datos:
            activo = datos.get("activo", True if actual is None else bool(actual["activo"]))
            if isinstance(activo, str):
                activo = activo.strip().lower() in ("1", "true", "activo", "activa")
            limpio["activo"] = 1 if activo else 0
        self.validar_extra(datos, limpio, errores, actual)
        if errores:
            raise ErrorNegocio(422, "validacion", "Revisa los campos marcados.", errores)
        self.validar_unicos(limpio, errores, actual)
        if errores:
            raise ErrorNegocio(409, "duplicado", "Ya existe un registro con esos datos.", errores)
        return limpio

    def validar_extra(self, datos, limpio, errores, actual):
        pass

    unicos = []  # (campo, mensaje)

    def validar_unicos(self, limpio, errores, actual):
        db = get_db()
        rid = actual["id"] if actual is not None else -1
        for campo, mensaje in self.unicos:
            if campo not in limpio:
                continue
            f = db.execute(f"SELECT id FROM {self.tabla} WHERE norm({campo}) = norm(?) AND id <> ?",
                           (limpio[campo], rid)).fetchone()
            if f:
                errores[campo] = mensaje

    # ------------------------------------------------------------- escritura
    def insertar(self, limpio):
        cols = [c for c in limpio if not c.startswith("_")]
        db = get_db()
        cur = db.execute(
            f"INSERT INTO {self.tabla} ({', '.join(cols)}) VALUES ({', '.join('?' * len(cols))})",
            [limpio[c] for c in cols])
        return cur.lastrowid

    def modificar(self, rid, limpio):
        cols = [c for c in limpio if not c.startswith("_")]
        get_db().execute(f"UPDATE {self.tabla} SET {', '.join(c + '=?' for c in cols)} WHERE id=?",
                         [limpio[c] for c in cols] + [rid])

    def crear(self, datos):
        limpio = self.validar(datos)
        self.antes_de_crear(limpio, datos)
        db = get_db()
        try:
            rid = self.insertar(limpio)
            self.despues_de_guardar(rid, limpio, datos, None)
            db.commit()
        except sqlite3.IntegrityError:
            db.rollback()
            raise ErrorNegocio(409, "duplicado", "Ya existe un registro con esos datos (código o nombre repetido).")
        return self.serializar(self.obtener(rid))

    def actualizar(self, rid, datos):
        actual = self.obtener(rid)
        limpio = self.validar(datos, actual)
        self.antes_de_actualizar(actual, limpio, datos)
        db = get_db()
        try:
            self.modificar(rid, limpio)
            self.despues_de_guardar(rid, limpio, datos, actual)
            db.commit()
        except sqlite3.IntegrityError:
            db.rollback()
            raise ErrorNegocio(409, "duplicado", "Ya existe un registro con esos datos (código o nombre repetido).")
        return self.serializar(self.obtener(rid))

    def cambiar_estado(self, rid, activo):
        actual = self.obtener(rid)
        activo = bool(activo)
        self.antes_de_estado(actual, activo)
        db = get_db()
        db.execute(f"UPDATE {self.tabla} SET activo=? WHERE id=?", (1 if activo else 0, rid))
        db.commit()
        return self.serializar(self.obtener(rid))

    def eliminar(self, rid):
        actual = self.obtener(rid)
        self.antes_de_eliminar(actual)
        db = get_db()
        try:
            db.execute(f"DELETE FROM {self.tabla} WHERE id=?", (rid,))
            db.commit()
        except sqlite3.IntegrityError:
            db.rollback()
            raise ErrorNegocio(409, "con_relaciones",
                               "No se puede eliminar porque tiene registros relacionados. Puedes desactivarlo.")

    # hooks de reglas de negocio
    def antes_de_crear(self, limpio, datos):
        pass

    def antes_de_actualizar(self, actual, limpio, datos):
        pass

    def despues_de_guardar(self, rid, limpio, datos, actual):
        pass

    def antes_de_estado(self, actual, activo):
        pass

    def antes_de_eliminar(self, actual):
        pass

    # ------------------------------------------------------------ exportación
    def filtro_descripcion(self, args):
        partes = []
        if limpiar_texto(args.get("q")):
            partes.append(f"Búsqueda: «{limpiar_texto(args.get('q'))[:100]}»")
        if args.get("estado") in ("activo", "inactivo"):
            partes.append(f"Estado: {args.get('estado')}")
        return "; ".join(partes) or "Sin filtros"

    def exportar(self, args):
        db = get_db()
        where, params = self._where(args)
        cond = (" WHERE " + " AND ".join(where)) if where else ""
        filas = db.execute(
            f"SELECT {self.seleccion()} {self.desde()}{cond} ORDER BY {self._orden(args)} LIMIT ?",
            params + [self.max_export]).fetchall()
        wb = Workbook()
        ws = wb.active
        ws.title = self.hoja[:31]
        encabezados = [h for h, _ in self.export_cols]
        ws.append(encabezados)
        relleno = PatternFill("solid", fgColor="920000")
        for i in range(1, len(encabezados) + 1):
            c = ws.cell(row=1, column=i)
            c.font, c.fill = Font(bold=True, color="FFFFFF"), relleno
            c.alignment = Alignment(vertical="center", wrap_text=True)
        anchos = [len(h) + 2 for h in encabezados]
        for f in filas:
            d = self.serializar(f)
            fila = [_seguro_excel(self.valor_export(d, k)) for _, k in self.export_cols]
            ws.append(fila)
            for i, v in enumerate(fila):
                anchos[i] = min(60, max(anchos[i], len(str(v)) + 2))
        for i, a in enumerate(anchos, start=1):
            ws.column_dimensions[get_column_letter(i)].width = a
        ws.freeze_panes = "A2"
        meta = wb.create_sheet("Filtro aplicado")
        ahora = datetime.now(ZONA).strftime("%d-%m-%Y %H:%M")
        for par in [("Módulo", self.etiqueta), ("Filtro", self.filtro_descripcion(args)),
                    ("Delegación activa", f"{g.delegacion['codigo']} · {g.delegacion['nombre']}"),
                    ("Registros exportados", len(filas)), ("Generado por", g.user["nombre"]),
                    ("Fecha de generación", ahora)]:
            meta.append(par)
        meta.column_dimensions["A"].width = 24
        meta.column_dimensions["B"].width = 70
        buf = io.BytesIO()
        wb.save(buf)
        return buf.getvalue(), len(filas)

    max_export = 50000

    def valor_export(self, d, clave):
        v = d.get(clave, "")
        if isinstance(v, list):
            return ", ".join(str(x) for x in v)
        if isinstance(v, bool):
            return "Sí" if v else "No"
        return v


def g_cfg(clave):
    from flask import current_app
    return current_app.config[clave]


# --------------------------------------------------------------------- rutas
def registrar_recurso(app, r):
    base = f"/api/{r.slug}"
    pref = r.clave

    def vista_listar():
        return jsonify(r.listar(request.args))

    def vista_obtener(rid):
        return jsonify({"ok": True, "item": r.serializar(r.obtener(rid))})

    def vista_exportar():
        contenido, _n = r.exportar(request.args)
        nombre = f"{r.slug}_{datetime.now(ZONA).strftime('%Y%m%d_%H%M')}.xlsx"
        return Response(contenido, mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                        headers={"Content-Disposition": f'attachment; filename="{nombre}"',
                                 "Cache-Control": "no-store"})

    app.add_url_rule(base, f"{pref}_listar", requiere(f"{pref}:ver")(vista_listar), methods=["GET"])
    app.add_url_rule(f"{base}/<int:rid>", f"{pref}_obtener", requiere(f"{pref}:ver")(vista_obtener), methods=["GET"])
    app.add_url_rule(f"{base}/exportar", f"{pref}_exportar", requiere(f"{pref}:exportar")(vista_exportar), methods=["GET"])

    if r.solo_lectura:
        return

    def vista_crear():
        item = r.crear(request.get_json(silent=True))
        resp = jsonify({"ok": True, "item": item, "mensaje": f"{r.etiqueta}: registro creado correctamente."})
        resp.status_code = 201
        return resp

    def vista_actualizar(rid):
        item = r.actualizar(rid, request.get_json(silent=True))
        return jsonify({"ok": True, "item": item, "mensaje": "Los cambios se guardaron correctamente."})

    def vista_estado(rid):
        cuerpo = request.get_json(silent=True) or {}
        if "activo" not in cuerpo or not isinstance(cuerpo["activo"], bool):
            raise ErrorNegocio(422, "validacion", "Indica el estado deseado (activo: true o false).")
        item = r.cambiar_estado(rid, cuerpo["activo"])
        return jsonify({"ok": True, "item": item,
                        "mensaje": f"El registro quedó {item['estado'].lower()}."})

    def vista_eliminar(rid):
        r.eliminar(rid)
        return jsonify({"ok": True, "mensaje": "El registro se eliminó correctamente."})

    app.add_url_rule(base, f"{pref}_crear", requiere(f"{pref}:crear")(vista_crear), methods=["POST"])
    app.add_url_rule(f"{base}/<int:rid>", f"{pref}_actualizar", requiere(f"{pref}:editar")(vista_actualizar), methods=["PUT"])
    app.add_url_rule(f"{base}/<int:rid>/estado", f"{pref}_estado", requiere(f"{pref}:editar")(vista_estado), methods=["PATCH"])
    app.add_url_rule(f"{base}/<int:rid>", f"{pref}_eliminar", requiere(f"{pref}:eliminar")(vista_eliminar), methods=["DELETE"])


def manejar_error_negocio(e):
    return error_json(e.status, e.codigo, e.mensaje, e.campos)
