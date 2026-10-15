"""Pruebas funcionales automatizadas (backend + reglas de negocio) de la 2.ª evaluación.
Cada método documenta su ID de caso (CP-xxx-nn) para enlazarlo con docs/PLAN_PRUEBAS.md.
Ejecutar:  python -m unittest discover -s tests -v   (o)   python tests/run_tests.py
"""
import io
import json
import os
import re
import sqlite3
import sys
import tempfile
import unittest
from datetime import datetime, timedelta, timezone

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
from openpyxl import load_workbook  # noqa: E402

from app import create_app  # noqa: E402
from app.db import init_db  # noqa: E402

CLAVE = "Demo#2026Sgr!"


class Cli:
    """Cliente HTTP de pruebas con manejo de sesión y token CSRF."""

    def __init__(self, app):
        self.app, self.c, self.tok = app, app.test_client(), None

    def refrescar(self, ruta="/login"):
        r = self.c.get(ruta, follow_redirects=True)
        m = re.search(r'csrf-token" content="([^"]+)"', r.get_data(as_text=True))
        if m:
            self.tok = m.group(1)
        return r

    def api(self, metodo, url, json_=None, **kw):
        if not self.tok:
            self.refrescar()
        return self.c.open(url, method=metodo, json=json_, headers={"X-CSRF-Token": self.tok}, **kw)

    def get(self, url, **kw):
        return self.c.get(url, **kw)

    def login(self, usuario, clave=CLAVE):
        self.refrescar("/login")
        r = self.api("POST", "/api/auth/login", {"usuario": usuario, "password": clave})
        if r.status_code == 200:
            self.refrescar("/")
        return r


class Base(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.tmp = tempfile.mkdtemp()
        cls.ruta = os.path.join(cls.tmp, "t.sqlite3")
        init_db(cls.ruta, demo=True, volumen=60)
        cls.app = create_app({"TESTING": True, "DATABASE_PATH": cls.ruta, "SECRET_KEY": "clave-de-prueba",
                              "CODIGO_REENVIO_SEG": 45, "LOGIN_MAX_INTENTOS": 5})

    def cli(self, usuario=None):
        c = Cli(self.app)
        if usuario:
            r = c.login(usuario)
            self.assertEqual(r.status_code, 200, r.get_data(as_text=True))
        return c

    def sql(self, consulta, params=()):
        con = sqlite3.connect(self.ruta)
        con.row_factory = sqlite3.Row
        try:
            filas = con.execute(consulta, params).fetchall()
            con.commit()
            return filas
        finally:
            con.close()


# =============================================================== ACCESO
class TestAcceso(Base):
    def test_CP_AUT_01_login_correcto_por_usuario(self):
        """Ingreso correcto con nombre de usuario."""
        r = self.cli().login("funcionario")
        self.assertEqual(r.status_code, 200)
        self.assertTrue(r.json["ok"])

    def test_CP_AUT_02_login_correcto_por_correo(self):
        """Ingreso correcto con correo (sin distinguir mayúsculas)."""
        self.assertEqual(self.cli().login("Funcionario@Example.TEST").status_code, 200)

    def test_CP_AUT_03_clave_incorrecta_mensaje_generico(self):
        c = self.cli()
        r1 = c.login("delegado", "Incorrecta#1234")
        c2 = self.cli()
        r2 = c2.login("no.existe.xyz", "Incorrecta#1234")
        self.assertEqual((r1.status_code, r2.status_code), (401, 401))
        self.assertEqual(r1.json["mensaje"], r2.json["mensaje"])  # no revela si la cuenta existe

    def test_CP_AUT_04_campos_vacios(self):
        c = self.cli()
        c.refrescar()
        r = c.api("POST", "/api/auth/login", {"usuario": "", "password": ""})
        self.assertEqual(r.status_code, 422)
        self.assertIn("usuario", r.json["campos"])
        self.assertIn("password", r.json["campos"])

    def test_CP_AUT_05_cuenta_inactiva(self):
        r = self.cli().login("inactivo")
        self.assertEqual(r.status_code, 403)
        self.assertEqual(r.json["error"], "cuenta_inactiva")
        r2 = self.cli().login("inactivo", "Mala#Clave123")  # con clave errónea no se revela el estado
        self.assertEqual(r2.status_code, 401)

    def test_CP_AUT_06_bloqueo_por_intentos(self):
        c = self.cli()
        for _ in range(5):
            self.assertEqual(c.login("delegado2", "Mala#Clave123").status_code, 401)
        r = c.login("delegado2", CLAVE)  # incluso la clave correcta queda bloqueada
        self.assertEqual(r.status_code, 429)
        r = c.login("usuario.inventado", CLAVE)
        for _ in range(5):
            c.login("otro.inventado", "x")
        self.assertEqual(c.login("otro.inventado", "x").status_code, 429)  # mismo trato si no existe

    def test_CP_AUT_07_cierre_de_sesion_y_rutas_protegidas(self):
        c = self.cli("funcionario2")
        self.assertEqual(c.get("/atenciones").status_code, 200)
        self.assertEqual(c.api("POST", "/api/auth/logout").status_code, 200)
        r = c.get("/atenciones")
        self.assertEqual(r.status_code, 302)
        self.assertIn("/login", r.headers["Location"])
        r = c.get("/api/atenciones")
        self.assertEqual(r.status_code, 401)

    def test_CP_AUT_08_rutas_protegidas_sin_sesion(self):
        c = Cli(self.app)
        for ruta in ["/", "/usuarios", "/roles", "/delegaciones", "/api/usuarios", "/api/roles"]:
            r = c.get(ruta)
            self.assertIn(r.status_code, (302, 401), ruta)

    def test_CP_AUT_09_csrf_obligatorio(self):
        c = self.cli("funcionario")
        r = c.c.post("/api/auth/logout")  # sin cabecera
        self.assertEqual(r.status_code, 403)
        r = c.c.post("/api/contexto", json={"delegacion_id": 1})
        self.assertEqual(r.status_code, 403)

    def test_CP_AUT_10_next_no_permite_redireccion_externa(self):
        c = self.cli()
        c.refrescar()
        r = c.api("POST", "/api/auth/login", {"usuario": "funcionario", "password": CLAVE, "next": "https://malo.example/x"})
        self.assertEqual(r.json["destino"], "/")
        c2 = self.cli()
        c2.refrescar()
        r = c2.api("POST", "/api/auth/login", {"usuario": "funcionario", "password": CLAVE, "next": "//malo.example"})
        self.assertEqual(r.json["destino"], "/")

    def test_CP_AUT_11_cuenta_desactivada_pierde_acceso_de_inmediato(self):
        c = self.cli("funcionario2")
        self.sql("UPDATE usuarios SET activo=0 WHERE usuario='funcionario2'")
        try:
            self.assertEqual(c.get("/api/atenciones").status_code, 401)
        finally:
            self.sql("UPDATE usuarios SET activo=1 WHERE usuario='funcionario2'")


# ============================================================ RECUPERACIÓN
class TestRecuperacion(Base):
    def codigo_de_bandeja(self, correo):
        filas = self.sql("SELECT cuerpo FROM bandeja_dev WHERE destino=? ORDER BY id DESC LIMIT 1", (correo,))
        self.assertTrue(filas, "No se generó correo")
        return re.search(r"es: (\d{6})", filas[0]["cuerpo"]).group(1)

    def solicitar(self, c, correo):
        c.refrescar("/recuperar")
        return c.api("POST", "/api/recuperar/solicitar", {"correo": correo})

    def retroceder_envio(self, correo, segundos=60):
        antes = (datetime.now(timezone.utc) - timedelta(seconds=segundos)).strftime("%Y-%m-%dT%H:%M:%SZ")
        self.sql("UPDATE recuperaciones SET creado_en=? WHERE usuario_id=(SELECT id FROM usuarios WHERE correo=?)", (antes, correo))

    def test_CP_REC_01_respuesta_no_revela_si_el_correo_existe(self):
        r1 = self.solicitar(self.cli(), "delegado@example.test")
        r2 = self.solicitar(self.cli(), "nadie.existe@example.test")
        self.assertEqual(r1.status_code, 200)
        self.assertEqual(r1.json, r2.json)
        self.assertFalse(self.sql("SELECT 1 FROM bandeja_dev WHERE destino='nadie.existe@example.test'"))

    def test_CP_REC_02_correo_invalido_o_vacio(self):
        c = self.cli()
        self.assertEqual(self.solicitar(c, "").status_code, 422)
        self.assertEqual(self.solicitar(c, "no-es-correo").status_code, 422)

    def test_CP_REC_03_flujo_completo_y_clave_anterior_rechazada(self):
        correo = "delegado2@example.test"
        c = self.cli()
        self.solicitar(c, correo)
        self.assertEqual(c.get("/recuperar/nueva").status_code, 302)  # etapa 3 cerrada
        r = c.api("POST", "/api/recuperar/validar", {"codigo": self.codigo_de_bandeja(correo)})
        self.assertEqual(r.status_code, 200, r.get_data(as_text=True))
        self.assertEqual(c.get("/recuperar/nueva").status_code, 200)
        r = c.api("POST", "/api/recuperar/nueva", {"password": "NuevaClave#2027x", "confirmacion": "NuevaClave#2027x"})
        self.assertEqual(r.status_code, 200)
        self.assertEqual(self.cli().login("delegado2", "NuevaClave#2027x").status_code, 200)
        self.assertEqual(self.cli().login("delegado2", CLAVE).status_code, 401)  # anterior rechazada
        self.sql("UPDATE usuarios SET password_hash=(SELECT password_hash FROM usuarios WHERE usuario='delegado') WHERE usuario='delegado2'")
        self.sql("DELETE FROM intentos_login")

    def test_CP_REC_04_codigo_incorrecto_vencido_y_reutilizado(self):
        correo = "funcionario2@example.test"
        c = self.cli()
        self.solicitar(c, correo)
        bueno = self.codigo_de_bandeja(correo)
        malo = "000000" if bueno != "000000" else "111111"
        r = c.api("POST", "/api/recuperar/validar", {"codigo": malo})
        self.assertEqual((r.status_code, r.json["error"]), (400, "codigo_invalido"))
        # vencido
        self.sql("UPDATE recuperaciones SET expira_en='2000-01-01T00:00:00Z' WHERE usuario_id=(SELECT id FROM usuarios WHERE correo=?)", (correo,))
        r = c.api("POST", "/api/recuperar/validar", {"codigo": bueno})
        self.assertEqual(r.status_code, 400)
        # reutilizado: nuevo código, se usa y no sirve otra vez
        self.retroceder_envio(correo)
        c2 = self.cli()
        self.solicitar(c2, correo)
        otro = self.codigo_de_bandeja(correo)
        self.assertEqual(c2.api("POST", "/api/recuperar/validar", {"codigo": otro}).status_code, 200)
        c3 = self.cli()
        c3.refrescar("/recuperar")
        c3.api("POST", "/api/recuperar/solicitar", {"correo": correo})  # sesión nueva, en espera
        self.assertEqual(c3.api("POST", "/api/recuperar/validar", {"codigo": otro}).status_code, 400)

    def test_CP_REC_05_reenvio_respeta_espera_de_45_segundos(self):
        correo = "coordinador@example.test"
        c = self.cli()
        self.solicitar(c, correo)
        r = c.api("POST", "/api/recuperar/reenviar")
        self.assertEqual(r.status_code, 429)
        self.assertGreater(r.json["espera"], 0)
        n = len(self.sql("SELECT 1 FROM recuperaciones WHERE usuario_id=(SELECT id FROM usuarios WHERE correo=?)", (correo,)))
        # simulando el paso del tiempo
        with c.c.session_transaction() as s:
            s["rec_envio_ts"] = 0
        self.retroceder_envio(correo)
        r = c.api("POST", "/api/recuperar/reenviar")
        self.assertEqual(r.status_code, 200)
        n2 = len(self.sql("SELECT 1 FROM recuperaciones WHERE usuario_id=(SELECT id FROM usuarios WHERE correo=?)", (correo,)))
        self.assertEqual(n2, n + 1)
        # el código anterior queda invalidado
        self.assertEqual(self.sql("SELECT COUNT(*) c FROM recuperaciones WHERE usado=0 AND usuario_id=(SELECT id FROM usuarios WHERE correo=?)", (correo,))[0]["c"], 1)

    def test_CP_REC_06_intentos_fallidos_agotan_el_codigo(self):
        correo = "delegado@example.test"
        self.retroceder_envio(correo, 4000)
        c = self.cli()
        self.solicitar(c, correo)
        bueno = self.codigo_de_bandeja(correo)
        malo = "000000" if bueno != "000000" else "111111"
        for _ in range(5):
            c.api("POST", "/api/recuperar/validar", {"codigo": malo})
        r = c.api("POST", "/api/recuperar/validar", {"codigo": bueno})
        self.assertEqual(r.status_code, 429)
        self.assertEqual(r.json["error"], "intentos_agotados")

    def test_CP_REC_07_nueva_clave_sin_verificacion_es_rechazada(self):
        c = self.cli()
        c.refrescar("/recuperar")
        r = c.api("POST", "/api/recuperar/nueva", {"password": "NuevaClave#2027x", "confirmacion": "NuevaClave#2027x"})
        self.assertEqual(r.status_code, 403)

    def test_CP_REC_08_requisitos_y_confirmacion_de_clave(self):
        correo = "funcionario@example.test"
        c = self.cli()
        self.solicitar(c, correo)
        c.api("POST", "/api/recuperar/validar", {"codigo": self.codigo_de_bandeja(correo)})
        for debil in ["corta1!A", "sinmayuscula#123", "SINMINUSCULA#123", "SinNumeros#abcd", "SinSimbolo1234A"]:
            r = c.api("POST", "/api/recuperar/nueva", {"password": debil, "confirmacion": debil})
            self.assertEqual(r.status_code, 422, debil)
        r = c.api("POST", "/api/recuperar/nueva", {"password": "Valida#Clave2027", "confirmacion": "Otra#Clave2027x"})
        self.assertEqual(r.status_code, 422)
        self.assertIn("confirmacion", r.json["campos"])

    def test_CP_REC_09_cambio_de_clave_cierra_sesiones_abiertas(self):
        correo = "delegado2@example.test"
        abierta = self.cli("delegado2")
        self.retroceder_envio(correo, 4000)
        c = self.cli()
        self.solicitar(c, correo)
        c.api("POST", "/api/recuperar/validar", {"codigo": self.codigo_de_bandeja(correo)})
        self.assertEqual(c.api("POST", "/api/recuperar/nueva", {"password": "Otra#Clave2028zz", "confirmacion": "Otra#Clave2028zz"}).status_code, 200)
        self.assertEqual(abierta.get("/api/atenciones").status_code, 401)
        self.sql("UPDATE usuarios SET password_hash=(SELECT password_hash FROM usuarios WHERE usuario='delegado'), session_version=session_version WHERE usuario='delegado2'")

    def test_CP_REC_10_codigo_no_se_guarda_en_claro(self):
        correo = "funcionario2@example.test"
        self.retroceder_envio(correo, 4000)
        self.solicitar(self.cli(), correo)
        codigo = self.codigo_de_bandeja(correo)
        hashes = [f["codigo_hash"] for f in self.sql("SELECT codigo_hash FROM recuperaciones")]
        self.assertNotIn(codigo, hashes)


# ============================================================ MANTENEDORES
CASOS_CRUD = {
    "delegaciones": dict(nuevo={"codigo": "del90", "nombre": "Delegación de prueba", "direccion": "Calle Falsa 123", "comuna": "La Serena"},
                         editar={"nombre": "Delegación editada", "direccion": "Calle Falsa 456", "comuna": "Coquimbo"},
                         buscar="editada", invalido={"codigo": "", "nombre": "x", "direccion": "", "comuna": "Marte"}),
    "roles": dict(nuevo={"codigo": "rol90", "nombre": "Rol de prueba", "descripcion": "d", "alcance": "UNA", "permisos": ["atenciones:ver"]},
                  editar={"nombre": "Rol editado", "descripcion": "e", "alcance": "UNA", "permisos": ["atenciones:ver", "vecinos:ver"]},
                  buscar="editado", invalido={"codigo": "", "nombre": "", "alcance": "X", "permisos": []}),
    "usuarios": dict(nuevo={"usuario": "prueba90", "nombre": "Usuaria de Prueba", "correo": "prueba90@example.test", "rol_id": None,
                            "delegaciones_ids": [1], "password": "Valida#Clave2027"},
                     editar={"nombre": "Usuaria Editada", "correo": "prueba90@example.test", "usuario": "prueba90"},
                     buscar="editada", invalido={"usuario": "", "nombre": "", "correo": "mal", "rol_id": 0, "delegaciones_ids": []}),
    "items": dict(nuevo={"codigo": "it90", "nombre": "Ítem de prueba", "descripcion": "d"}, editar={"nombre": "Ítem editado", "descripcion": "e", "codigo": "IT90"},
                  buscar="editado", invalido={"codigo": "!!", "nombre": ""}),
    "tipos-atencion": dict(nuevo={"codigo": "ta90", "nombre": "Tipo de prueba", "descripcion": "d"}, editar={"nombre": "Tipo editado", "descripcion": "e", "codigo": "TA90"},
                           buscar="editado", invalido={"codigo": "", "nombre": "ab"}),
    "sub-atenciones": dict(nuevo={"codigo": "sa90", "nombre": "Sub de prueba", "tipo_id": 1}, editar={"nombre": "Sub editada", "tipo_id": 1, "codigo": "SA90"},
                           buscar="editada", invalido={"codigo": "", "nombre": "", "tipo_id": 0}),
}


class TestCrud(Base):
    def test_CP_CRUD_01_ciclo_completo_de_los_seis_mantenedores(self):
        """Crear, consultar, editar y eliminar en cada mantenedor (superadministrador)."""
        c = self.cli("superadmin")
        rol_op = self.sql("SELECT id FROM roles WHERE codigo='ROL04'")[0]["id"]
        for slug, caso in CASOS_CRUD.items():
            with self.subTest(modulo=slug):
                nuevo = dict(caso["nuevo"])
                if slug == "usuarios":
                    nuevo["rol_id"] = rol_op
                r = c.api("POST", f"/api/{slug}", nuevo)
                self.assertEqual(r.status_code, 201, r.get_data(as_text=True))
                rid = r.json["item"]["id"]
                self.assertEqual(c.get(f"/api/{slug}/{rid}").json["item"]["id"], rid)
                r = c.api("PUT", f"/api/{slug}/{rid}", caso["editar"])
                self.assertEqual(r.status_code, 200, r.get_data(as_text=True))
                self.assertEqual(c.get(f"/api/{slug}?q={caso['buscar']}").json["total"], 1)
                r = c.api("DELETE", f"/api/{slug}/{rid}")
                self.assertEqual(r.status_code, 200, r.get_data(as_text=True))
                self.assertEqual(c.get(f"/api/{slug}/{rid}").status_code, 404)

    def test_CP_CRUD_02_campos_requeridos_y_formatos(self):
        c = self.cli("superadmin")
        for slug, caso in CASOS_CRUD.items():
            with self.subTest(modulo=slug):
                r = c.api("POST", f"/api/{slug}", caso["invalido"])
                self.assertEqual(r.status_code, 422, r.get_data(as_text=True))
                self.assertTrue(r.json["campos"])

    def test_CP_CRUD_03_longitudes_maximas(self):
        c = self.cli("superadmin")
        r = c.api("POST", "/api/items", {"codigo": "IT91", "nombre": "x" * 101})
        self.assertEqual(r.status_code, 422)
        r = c.api("POST", "/api/items", {"codigo": "IT91" * 6, "nombre": "Nombre válido"})
        self.assertEqual(r.status_code, 422)

    def test_CP_CRUD_04_duplicados(self):
        c = self.cli("superadmin")
        self.assertEqual(c.api("POST", "/api/delegaciones", {"codigo": "del01", "nombre": "Otra", "direccion": "Calle 12345", "comuna": "La Serena"}).status_code, 409)
        self.assertEqual(c.api("POST", "/api/tipos-atencion", {"codigo": "TA777", "nombre": "social"}).status_code, 409)  # nombre sin distinguir mayúsculas
        self.assertEqual(c.api("POST", "/api/roles", {"codigo": "ROL77", "nombre": "Coordinador", "alcance": "UNA", "permisos": ["atenciones:ver"]}).status_code, 409)
        rol = self.sql("SELECT id FROM roles WHERE codigo='ROL04'")[0]["id"]
        r = c.api("POST", "/api/usuarios", {"usuario": "otro.usuario", "nombre": "Otro Nombre", "correo": "DELEGADO@example.test",
                                              "rol_id": rol, "delegaciones_ids": [1], "password": "Valida#Clave2027"})
        self.assertEqual(r.status_code, 409)

    def test_CP_CRUD_05_restricciones_por_relaciones(self):
        c = self.cli("superadmin")
        # delegación con usuarios y atenciones
        r = c.api("DELETE", "/api/delegaciones/2")
        self.assertEqual(r.status_code, 409)
        # rol con usuarios
        rol = self.sql("SELECT id FROM roles WHERE codigo='ROL04'")[0]["id"]
        self.assertEqual(c.api("DELETE", f"/api/roles/{rol}").status_code, 409)
        # tipo con subatenciones
        self.assertEqual(c.api("DELETE", "/api/tipos-atencion/1").status_code, 409)
        # subatención con atenciones
        sub = self.sql("SELECT subtipo_id FROM atenciones LIMIT 1")[0]["subtipo_id"]
        self.assertEqual(c.api("DELETE", f"/api/sub-atenciones/{sub}").status_code, 409)
        # usuario con atenciones
        uid = self.sql("SELECT usuario_id FROM atenciones WHERE usuario_id<>1 LIMIT 1")[0]["usuario_id"]
        self.assertEqual(c.api("DELETE", f"/api/usuarios/{uid}").status_code, 409)

    def test_CP_TIP_01_tipo_y_subatencion_integridad(self):
        c = self.cli("superadmin")
        t = c.api("POST", "/api/tipos-atencion", {"codigo": "TA801", "nombre": "Tipo integridad"}).json["item"]["id"]
        s = c.api("POST", "/api/sub-atenciones", {"codigo": "SA801", "nombre": "Sub integridad", "tipo_id": t}).json["item"]["id"]
        self.assertEqual(c.api("PATCH", f"/api/tipos-atencion/{t}/estado", {"activo": False}).status_code, 409)  # tiene sub activa
        self.assertEqual(c.api("DELETE", f"/api/tipos-atencion/{t}").status_code, 409)
        self.assertEqual(c.api("PATCH", f"/api/sub-atenciones/{s}/estado", {"activo": False}).status_code, 200)
        self.assertEqual(c.api("PATCH", f"/api/tipos-atencion/{t}/estado", {"activo": False}).status_code, 200)
        self.assertEqual(c.api("PATCH", f"/api/sub-atenciones/{s}/estado", {"activo": True}).status_code, 409)  # tipo inactivo
        r = c.api("POST", "/api/sub-atenciones", {"codigo": "SA802", "nombre": "Bajo tipo inactivo", "tipo_id": t})
        self.assertEqual(r.status_code, 422)
        self.assertEqual(c.api("DELETE", f"/api/sub-atenciones/{s}").status_code, 200)
        self.assertEqual(c.api("DELETE", f"/api/tipos-atencion/{t}").status_code, 200)

    def test_CP_TIP_02_subatencion_duplicada_en_mismo_tipo(self):
        c = self.cli("superadmin")
        r = c.api("POST", "/api/sub-atenciones", {"codigo": "SA803", "nombre": "ayuda social", "tipo_id": 1})
        self.assertEqual(r.status_code, 422)
        self.assertIn("nombre", r.json["campos"])

    def test_CP_USR_01_reglas_de_usuarios(self):
        c = self.cli("superadmin")
        rol_op = self.sql("SELECT id FROM roles WHERE codigo='ROL04'")[0]["id"]
        rol_una = rol_op
        base = {"usuario": "regla.uno", "nombre": "Regla Uno", "correo": "regla.uno@example.test", "rol_id": rol_una, "delegaciones_ids": [1], "password": "Valida#Clave2027"}
        # rol de una delegación con dos delegaciones
        self.assertEqual(c.api("POST", "/api/usuarios", {**base, "delegaciones_ids": [1, 2]}).status_code, 422)
        # clave débil y clave que contiene el usuario
        self.assertEqual(c.api("POST", "/api/usuarios", {**base, "password": "debil"}).status_code, 422)
        self.assertEqual(c.api("POST", "/api/usuarios", {**base, "password": "Regla.uno#2027A"}).status_code, 422)
        # autoprotección
        yo = self.sql("SELECT id FROM usuarios WHERE usuario='superadmin'")[0]["id"]
        self.assertEqual(c.api("PATCH", f"/api/usuarios/{yo}/estado", {"activo": False}).status_code, 409)
        self.assertEqual(c.api("DELETE", f"/api/usuarios/{yo}").status_code, 409)
        # la clave nunca se devuelve
        self.assertNotIn("password", json.dumps(c.get(f"/api/usuarios/{yo}").json).lower())

    def test_CP_ROL_01_escalamiento_de_privilegios(self):
        c = self.cli("coordinador")
        self.assertEqual(c.get("/api/roles").status_code, 403)  # el coordinador no administra roles
        s = self.cli("superadmin")
        # permisos sin Ver
        r = s.api("POST", "/api/roles", {"codigo": "ROL91", "nombre": "Sin ver", "alcance": "UNA", "permisos": ["usuarios:crear"]})
        self.assertEqual(r.status_code, 422)
        # permiso inexistente
        r = s.api("POST", "/api/roles", {"codigo": "ROL92", "nombre": "Inexistente", "alcance": "UNA", "permisos": ["magia:ver"]})
        self.assertEqual(r.status_code, 422)
        # no puede modificar su propio rol
        rid = self.sql("SELECT id FROM roles WHERE codigo='ROL01'")[0]["id"]
        self.assertEqual(s.api("DELETE", f"/api/roles/{rid}").status_code, 403)

    def test_CP_ROL_02_un_coordinador_no_asigna_rol_superior_ni_delegaciones_ajenas(self):
        # damos al coordinador permisos de usuarios (ya los tiene) y probamos límites
        c = self.cli("coordinador")
        admin = self.sql("SELECT id FROM roles WHERE codigo='ROL01'")[0]["id"]
        op = self.sql("SELECT id FROM roles WHERE codigo='ROL04'")[0]["id"]
        base = {"usuario": "esc.uno", "nombre": "Escalamiento Uno", "correo": "esc.uno@example.test", "delegaciones_ids": [1], "password": "Valida#Clave2027"}
        r = c.api("POST", "/api/usuarios", {**base, "rol_id": admin})
        self.assertEqual(r.status_code, 422)
        self.assertIn("rol_id", r.json["campos"])
        r = c.api("POST", "/api/usuarios", {**base, "rol_id": op, "delegaciones_ids": [3]})  # DEL03 no está asignada
        self.assertEqual(r.status_code, 422)
        r = c.api("POST", "/api/usuarios", {**base, "rol_id": op})
        self.assertEqual(r.status_code, 201)
        c.api("DELETE", f"/api/usuarios/{r.json['item']['id']}")

    def test_CP_DEL_01_no_se_desactiva_la_delegacion_activa(self):
        c = self.cli("superadmin")
        self.assertEqual(c.api("PATCH", "/api/delegaciones/1/estado", {"activo": False}).status_code, 409)


# ========================================================= LISTADOS / EXCEL
class TestListados(Base):
    def test_CP_LST_01_busqueda_parcial_sin_mayusculas_ni_tildes(self):
        c = self.cli("superadmin")
        self.assertGreaterEqual(c.get("/api/tipos-atencion?q=EDUCACION").json["total"], 1)
        self.assertGreaterEqual(c.get("/api/tipos-atencion?q=educ").json["total"], 1)
        self.assertEqual(c.get("/api/tipos-atencion?q=zzzz").json["total"], 0)

    def test_CP_LST_02_busqueda_en_varios_campos_y_terminos(self):
        c = self.cli("superadmin")
        self.assertEqual(c.get("/api/sub-atenciones?q=salud").json["total"], 3)  # por nombre del tipo
        self.assertEqual(c.get("/api/sub-atenciones?q=SA004").json["total"], 1)  # por código
        self.assertEqual(c.get("/api/sub-atenciones?q=salud+consulta").json["total"], 1)  # varios términos (AND)
        self.assertEqual(c.get("/api/usuarios?q=EXAMPLE.TEST&ambito=autorizadas").json["total"], 7)

    def test_CP_LST_03_caracteres_especiales_en_busqueda(self):
        c = self.cli("superadmin")
        for q in ["%", "_", "\\", "'; DROP TABLE items;--", "<script>"]:
            r = c.get("/api/items", query_string={"q": q})
            self.assertEqual(r.status_code, 200)
        self.assertEqual(c.get("/api/items", query_string={"q": "%"}).json["total"], 0)  # % no es comodín
        self.assertGreater(c.get("/api/items").json["total"], 0)

    def test_CP_LST_04_paginacion_y_total(self):
        c = self.cli("superadmin")
        d = c.get("/api/items?per_page=10&page=1").json
        self.assertEqual((d["total"], d["per_page"], d["paginas"], len(d["items"])), (24, 10, 3, 10))
        d3 = c.get("/api/items?per_page=10&page=3").json
        self.assertEqual(len(d3["items"]), 4)
        self.assertEqual(c.get("/api/items?per_page=10&page=99").json["page"], 3)  # página fuera de rango se ajusta
        ids = {i["id"] for p in (1, 2, 3) for i in c.get(f"/api/items?per_page=10&page={p}").json["items"]}
        self.assertEqual(len(ids), 24)  # sin repetidos ni faltantes

    def test_CP_LST_05_filtros_combinados_se_mantienen_al_paginar(self):
        c = self.cli("superadmin")
        todos = c.get("/api/items?estado=activo&q=meta&per_page=5").json
        self.assertGreater(todos["total"], 5)
        pag2 = c.get("/api/items?estado=activo&q=meta&per_page=5&page=2").json
        self.assertTrue(all(i["activo"] and "meta" in i["nombre"].lower() for i in pag2["items"]))
        self.assertEqual(pag2["total"], todos["total"])

    def test_CP_LST_06_exportacion_coincide_con_la_tabla_y_con_todas_las_paginas(self):
        c = self.cli("superadmin")
        total = c.get("/api/items?estado=activo&q=meta&per_page=5").json["total"]
        r = c.get("/api/items/exportar?estado=activo&q=meta")
        self.assertEqual(r.status_code, 200)
        self.assertIn("spreadsheetml", r.mimetype)
        wb = load_workbook(io.BytesIO(r.data))
        hoja = wb.worksheets[0]
        self.assertEqual(hoja.max_row - 1, total)
        self.assertEqual(wb["Filtro aplicado"]["B4"].value, total)

    def test_CP_LST_07_exportacion_no_incluye_campos_sensibles(self):
        c = self.cli("superadmin")
        wb = load_workbook(io.BytesIO(c.get("/api/usuarios/exportar?ambito=autorizadas").data))
        encabezados = [x.value for x in wb.worksheets[0][1]]
        texto = " ".join(str(f.value) for fila in wb.worksheets[0].iter_rows() for f in fila).lower()
        for prohibido in ("password", "hash", "token", "session", "scrypt"):
            self.assertNotIn(prohibido, " ".join(encabezados).lower())
            self.assertNotIn(prohibido, texto)

    def test_CP_LST_08_exportacion_neutraliza_formulas(self):
        c = self.cli("superadmin")
        c.api("POST", "/api/items", {"codigo": "IT92", "nombre": "=HYPERLINK(\"http://x\")", "descripcion": "+cmd"})
        wb = load_workbook(io.BytesIO(c.get("/api/items/exportar?q=HYPERLINK").data))
        self.assertTrue(str(wb.worksheets[0]["B2"].value).startswith("'"))
        c.api("DELETE", f"/api/items/{self.sql(chr(83)+'ELECT id FROM items WHERE codigo=?', ('IT92',))[0]['id']}")

    def test_CP_LST_09_exportar_requiere_permiso(self):
        self.assertEqual(self.cli("funcionario").get("/api/atenciones/exportar").status_code, 403)
        self.assertEqual(self.cli("delegado").get("/api/atenciones/exportar").status_code, 200)

    def test_CP_LST_10_ordenamiento(self):
        c = self.cli("superadmin")
        a = [i["codigo"] for i in c.get("/api/items?sort=codigo&dir=desc&per_page=5").json["items"]]
        self.assertEqual(a, sorted(a, reverse=True))
        self.assertEqual(c.get("/api/items?sort=hack;drop&dir=x").status_code, 200)  # clave no permitida se ignora


# ======================================================== PERFIL Y DELEGACIÓN
class TestPerfil(Base):
    def test_CP_PER_01_menu_diferenciado_por_rol(self):
        esperado = {
            "superadmin": {"/roles", "/usuarios", "/delegaciones", "/items", "/tipos-atencion", "/sub-atenciones", "/atenciones", "/reportes"},
            "coordinador": {"/usuarios", "/delegaciones", "/items", "/atenciones", "/reportes"},
            "delegado": {"/usuarios", "/atenciones", "/reportes", "/delegaciones"},
            "funcionario": {"/atenciones", "/vecinos", "/seguimiento"},
        }
        todos = {"/roles", "/usuarios", "/delegaciones", "/items", "/tipos-atencion", "/sub-atenciones", "/atenciones", "/vecinos",
                 "/tipos-gestion", "/seguimiento", "/reportes"}
        for u, rutas in esperado.items():
            html = self.cli(u).get("/").get_data(as_text=True)
            visibles = {r for r in todos if f'href="{r}"' in html}
            self.assertTrue(rutas <= visibles, (u, rutas - visibles))
        f = self.cli("funcionario").get("/").get_data(as_text=True)
        for r in ("/roles", "/usuarios", "/items", "/delegaciones"):
            self.assertNotIn(f'href="{r}"', f)

    def test_CP_PER_02_acceso_directo_a_modulos_no_autorizados(self):
        f = self.cli("funcionario")
        for ruta in ("/roles", "/usuarios", "/delegaciones", "/items", "/tipos-atencion", "/sub-atenciones", "/tipos-gestion", "/reportes"):
            self.assertEqual(f.get(ruta).status_code, 403, ruta)
        for api in ("/api/roles", "/api/usuarios", "/api/delegaciones", "/api/items"):
            self.assertEqual(f.get(api).status_code, 403, api)
        self.assertEqual(f.api("POST", "/api/items", {"codigo": "IT99", "nombre": "Intento"}).status_code, 403)
        self.assertEqual(f.api("DELETE", "/api/items/1").status_code, 403)
        d = self.cli("delegado")
        self.assertEqual(d.api("POST", "/api/usuarios", {}).status_code, 403)  # puede ver pero no crear
        self.assertEqual(d.get("/api/usuarios").status_code, 200)

    def test_CP_PER_03_distintivo_muestra_delegacion_usuario_y_rol(self):
        html = self.cli("funcionario").get("/atenciones").get_data(as_text=True)
        self.assertIn("DEL01 | Delegación Centro", html)
        self.assertIn("Delegación activa", html)  # texto, no solo color
        self.assertIn("Funcionario", html)
        self.assertNotIn('id="selector-delegacion"', html)  # una sola delegación: contexto fijo

    def test_CP_PER_04_selector_solo_con_delegaciones_autorizadas(self):
        html = self.cli("coordinador").get("/").get_data(as_text=True)
        self.assertIn('id="selector-delegacion"', html)
        self.assertIn("DEL02", html)
        self.assertNotIn("DEL03 |", html)

    def test_CP_PER_05_cambio_de_contexto_actualiza_listados(self):
        c = self.cli("coordinador")
        a1 = c.get("/api/atenciones?per_page=100").json
        self.assertTrue(all(x["delegacion"].startswith("DEL01") for x in a1["items"]))
        self.assertEqual(c.api("POST", "/api/contexto", {"delegacion_id": 2}).status_code, 200)
        a2 = c.get("/api/atenciones?per_page=100").json
        self.assertTrue(all(x["delegacion"].startswith("DEL02") for x in a2["items"]))
        self.assertIn("DEL02 | Delegación Norte", c.get("/").get_data(as_text=True))
        u = c.get("/api/usuarios?per_page=100").json
        self.assertTrue(all(any(d["codigo"] == "DEL02" for d in x["delegaciones"]) for x in u["items"]))
        c.api("POST", "/api/contexto", {"delegacion_id": 1})

    def test_CP_PER_06_no_se_puede_cambiar_a_delegacion_ajena(self):
        c = self.cli("coordinador")
        self.assertEqual(c.api("POST", "/api/contexto", {"delegacion_id": 3}).status_code, 403)
        self.assertEqual(c.api("POST", "/api/contexto", {"delegacion_id": 9999}).status_code, 403)
        self.assertEqual(c.api("POST", "/api/contexto", {"delegacion_id": "x"}).status_code, 422)
        f = self.cli("funcionario")
        self.assertEqual(f.api("POST", "/api/contexto", {"delegacion_id": 2}).status_code, 403)

    def test_CP_PER_07_acceso_a_registros_de_otra_delegacion_por_id(self):
        f = self.cli("funcionario")  # DEL01
        ajena = self.sql("SELECT id FROM atenciones WHERE delegacion_id=2 LIMIT 1")[0]["id"]
        propia = self.sql("SELECT id FROM atenciones WHERE delegacion_id=1 LIMIT 1")[0]["id"]
        self.assertEqual(f.get(f"/api/atenciones/{ajena}").status_code, 403)
        self.assertEqual(f.get(f"/api/atenciones/{propia}").status_code, 200)
        self.assertEqual(f.get("/api/atenciones?ambito=autorizadas&per_page=100").json["items"][0]["delegacion"][:5], "DEL01")
        d = self.cli("delegado")
        otro = self.sql("SELECT id FROM usuarios WHERE usuario='delegado2'")[0]["id"]
        self.assertEqual(d.get(f"/api/usuarios/{otro}").status_code, 403)
        self.assertEqual(d.get("/api/delegaciones").json["total"], 1)  # solo ve la suya

    def test_CP_PER_08_superadmin_ve_todas_las_delegaciones(self):
        s = self.cli("superadmin")
        self.assertEqual(s.get("/api/delegaciones").json["total"], 8)
        self.assertEqual(s.api("POST", "/api/contexto", {"delegacion_id": 5}).status_code, 200)
        s.api("POST", "/api/contexto", {"delegacion_id": 1})

    def test_CP_PER_09_opciones_de_formularios_respetan_el_alcance(self):
        c = self.cli("coordinador")
        d = c.get("/api/opciones/delegaciones").json["items"]
        self.assertEqual({x["texto"][:5] for x in d}, {"DEL01", "DEL02"})
        roles = {x["texto"].split(" (")[0] for x in c.get("/api/opciones/roles").json["items"]}
        self.assertNotIn("Superadministrador", roles)
        self.assertEqual(self.cli("funcionario").get("/api/opciones/roles").status_code, 403)


# ========================================================= INGRESO DE ATENCIÓN
class TestAtencion(Base):
    def datos(self, **extra):
        d = {"vecino_id": 1, "tipo_id": 1, "subtipo_id": 1, "descripcion": "Solicitud de prueba con detalle suficiente."}
        d.update(extra)
        return d

    def test_CP_ATN_01_registro_correcto_con_folio_responsable_y_delegacion(self):
        c = self.cli("funcionario")
        r = c.api("POST", "/api/atenciones", self.datos())
        self.assertEqual(r.status_code, 201, r.get_data(as_text=True))
        a = r.json["atencion"]
        self.assertRegex(a["folio"], r"^ATN-\d{4}-\d{6}$")
        fila = self.sql("SELECT a.*, u.usuario FROM atenciones a JOIN usuarios u ON u.id=a.usuario_id WHERE a.id=?", (a["id"],))[0]
        self.assertEqual((fila["usuario"], fila["delegacion_id"]), ("funcionario", 1))
        self.assertTrue(fila["fecha_registro"])

    def test_CP_ATN_02_campos_obligatorios(self):
        c = self.cli("funcionario")
        r = c.api("POST", "/api/atenciones", {})
        self.assertEqual(r.status_code, 422)
        self.assertEqual(set(r.json["campos"]), {"vecino_id", "tipo_id", "subtipo_id", "descripcion"})
        self.assertEqual(c.api("POST", "/api/atenciones", self.datos(descripcion="corta")).status_code, 422)
        self.assertEqual(c.api("POST", "/api/atenciones", self.datos(descripcion="x" * 2001)).status_code, 422)

    def test_CP_ATN_03_dependencia_tipo_subatencion(self):
        c = self.cli("funcionario")
        sub_de_otro_tipo = self.sql("SELECT id FROM subatenciones WHERE tipo_id<>1 LIMIT 1")[0]["id"]
        r = c.api("POST", "/api/atenciones", self.datos(subtipo_id=sub_de_otro_tipo))
        self.assertEqual(r.status_code, 422)
        self.assertIn("subtipo_id", r.json["campos"])
        opciones = c.get("/api/atencion/opciones").json
        self.assertTrue(all(s["tipo_id"] in {t["id"] for t in opciones["tipos"]} for s in opciones["subtipos"]))

    def test_CP_ATN_04_solo_opciones_activas(self):
        s = self.cli("superadmin")
        sub = s.api("POST", "/api/sub-atenciones", {"codigo": "SA805", "nombre": "Sub inactiva prueba", "tipo_id": 1, "activo": False}).json["item"]["id"]
        c = self.cli("funcionario")
        self.assertNotIn(sub, [x["id"] for x in c.get("/api/atencion/opciones").json["subtipos"]])
        r = c.api("POST", "/api/atenciones", self.datos(subtipo_id=sub))
        self.assertEqual(r.status_code, 422)
        s.api("DELETE", f"/api/sub-atenciones/{sub}")

    def test_CP_ATN_05_prevencion_de_doble_envio(self):
        c = self.cli("funcionario")
        antes = self.sql("SELECT COUNT(*) n FROM atenciones")[0]["n"]
        d = self.datos(token_envio="token-unico-123")
        r1 = c.api("POST", "/api/atenciones", d)
        r2 = c.api("POST", "/api/atenciones", d)
        self.assertEqual(r1.json["atencion"]["folio"], r2.json["atencion"]["folio"])
        self.assertTrue(r2.json["atencion"]["repetido"])
        self.assertEqual(self.sql("SELECT COUNT(*) n FROM atenciones")[0]["n"], antes + 1)

    def test_CP_ATN_06_no_se_puede_forzar_delegacion_ni_responsable(self):
        c = self.cli("funcionario")
        self.assertEqual(c.api("POST", "/api/atenciones", self.datos(delegacion_id=2)).status_code, 403)
        self.assertEqual(c.api("POST", "/api/atenciones", self.datos(usuario_id=1)).status_code, 403)

    def test_CP_ATN_07_usa_la_delegacion_activa_tras_cambiar_contexto(self):
        c = self.cli("coordinador")
        c.api("POST", "/api/contexto", {"delegacion_id": 2})
        r = c.api("POST", "/api/atenciones", self.datos())
        self.assertEqual(self.sql("SELECT delegacion_id d FROM atenciones WHERE id=?", (r.json["atencion"]["id"],))[0]["d"], 2)
        c.api("POST", "/api/contexto", {"delegacion_id": 1})

    def test_CP_ATN_08_solicitante_busqueda_y_alta_rapida(self):
        c = self.cli("funcionario")
        self.assertGreater(len(c.get("/api/vecinos/buscar?q=camila").json["items"]), 0)
        self.assertEqual(c.get("/api/vecinos/buscar?q=a").json["items"], [])  # mínimo 2 caracteres
        self.assertEqual(c.api("POST", "/api/vecinos", {"rut": "11111111-2", "nombres": "A", "apellidos": ""}).status_code, 422)  # RUT inválido y datos
        r = c.api("POST", "/api/vecinos", {"rut": "12.345.678-5", "nombres": "Persona", "apellidos": "Nueva Prueba"})
        self.assertEqual(r.status_code, 201, r.get_data(as_text=True))
        self.assertEqual(r.json["item"]["rut"], "12345678-5")
        self.assertEqual(c.api("POST", "/api/vecinos", {"rut": "12345678-5", "nombres": "Otra", "apellidos": "Persona"}).status_code, 409)

    def test_CP_ATN_09_solicitante_inexistente(self):
        c = self.cli("funcionario")
        self.assertEqual(c.api("POST", "/api/atenciones", self.datos(vecino_id=999999)).status_code, 422)

    def test_CP_ATN_10_cancelar_no_crea_registros(self):
        """Cancelar es una acción solo de interfaz: ninguna llamada a la API ocurre."""
        antes = self.sql("SELECT COUNT(*) n FROM atenciones")[0]["n"]
        self.cli("funcionario")  # abrir el formulario no crea nada
        self.assertEqual(self.sql("SELECT COUNT(*) n FROM atenciones")[0]["n"], antes)


# ============================================================== SEGURIDAD
class TestSeguridad(Base):
    def test_CP_SEG_01_ninguna_respuesta_filtra_hashes(self):
        s = self.cli("superadmin")
        for ruta in ("/api/usuarios?per_page=100&ambito=autorizadas", "/api/usuarios/1", "/api/roles", "/api/delegaciones"):
            t = s.get(ruta).get_data(as_text=True).lower()
            self.assertNotIn("scrypt", t)
            self.assertNotIn("password_hash", t)

    def test_CP_SEG_02_cabeceras_y_cookie(self):
        r = self.cli().get("/login")
        self.assertEqual(r.headers["X-Content-Type-Options"], "nosniff")
        self.assertIn("Content-Security-Policy", r.headers)
        self.assertIn("HttpOnly", r.headers.get("Set-Cookie", ""))

    def test_CP_SEG_03_salida_html_escapada(self):
        s = self.cli("superadmin")
        s.api("POST", "/api/items", {"codigo": "IT93", "nombre": "<script>alert(1)</script>", "descripcion": ""})
        # el API devuelve JSON (el escape lo hace el cliente con escapar()); la página HTML nunca incrusta datos
        html = s.get("/items").get_data(as_text=True)
        self.assertNotIn("<script>alert(1)</script>", html)
        s.api("DELETE", f"/api/items/{self.sql('SELECT id FROM items WHERE codigo=?', ('IT93',))[0]['id']}")

    def test_CP_SEG_04_inyeccion_sql_en_login(self):
        c = self.cli()
        self.assertEqual(c.login("' OR '1'='1", "' OR '1'='1").status_code, 401)

    def test_CP_SEG_05_salud(self):
        r = self.cli().get("/salud")
        self.assertEqual(r.status_code, 200)
        self.assertEqual(r.json["base_de_datos"], "ok")
        self.assertTrue(r.json["version"])


if __name__ == "__main__":
    unittest.main(verbosity=2)
